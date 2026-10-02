#!/usr/bin/env python3
"""S15：把 content.json 渲染成本地 docx（兜底成品）。

用法:
    python scripts/render_docx.py --dir "data/0001.《标题》"

红线落实:
- 真 Heading 1/2 样式（非加粗），否则大纲/飞书导入无结构
- 开头就有可点击的原视频链接
- 每张图都带「📍 hh:mm:ss」标注
- 中文字体显式设置 eastasia，否则显示方框
"""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
GRAY = RGBColor(0x8A, 0x90, 0x98)
INK = RGBColor(0x2B, 0x2F, 0x36)


def set_cjk(font_obj, name: str = "微软雅黑") -> None:
    """python-docx 必须同时设 ascii 与 eastasia，中文才不会变方框。

    对 style.font 和 run.font 都适用：先让 .name 生成 rFonts，
    再找到它补上 eastAsia。
    """
    font_obj.name = name
    el = font_obj._element
    rfonts = el.find(".//" + qn("w:rFonts"))
    if rfonts is not None:
        rfonts.set(qn("w:eastAsia"), name)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    p = paragraph._p
    r_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    hl = p.makeelement(qn("w:hyperlink"), {qn("r:id"): r_id})
    run = p.makeelement(qn("w:r"), {})
    rPr = p.makeelement(qn("w:rPr"), {})
    color = p.makeelement(qn("w:color"), {qn("w:val"): "0563C1"})
    underline = p.makeelement(qn("w:u"), {qn("w:val"): "single"})
    rPr.append(color)
    rPr.append(underline)
    run.append(rPr)
    t = p.makeelement(qn("w:t"), {})
    t.text = text
    run.append(t)
    hl.append(run)
    p.append(hl)


def caption(doc: Document, text: str, center: bool = True) -> None:
    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.size = Pt(9)
    r.font.color.rgb = GRAY
    set_cjk(r.font)


def picture(doc: Document, video_dir: Path, rel: str, cap: str) -> None:
    path = video_dir / rel
    if not path.exists():
        caption(doc, f"（缺图：{rel}）")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(15))
    caption(doc, f"📍 {cap}")


def main() -> None:
    ap = argparse.ArgumentParser(description="content.json → 本地 docx")
    ap.add_argument("--dir", required=True, help="视频文件夹")
    a = ap.parse_args()
    video_dir = (ROOT / a.dir).resolve() if not Path(a.dir).is_absolute() else Path(a.dir)
    c = json.loads((video_dir / "content.json").read_text(encoding="utf-8"))
    meta = c["meta"]

    doc = Document()
    for s in doc.styles:
        font_obj = getattr(s, "font", None)
        if font_obj is not None:
            try:
                set_cjk(font_obj)
            except Exception:
                pass
    for name, size, color in (("Heading 1", 20, INK), ("Heading 2", 15, INK),
                              ("Title", 26, INK)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.color.rgb = color
        set_cjk(st.font)

    doc.add_heading(f"{meta['num']}.《{meta['title']}》", level=0)

    p = doc.add_paragraph()
    r = p.add_run("原视频：")
    r.font.bold = True
    add_hyperlink(p, meta["url"], meta["url"])
    info = doc.add_paragraph()
    info.add_run(
        f"UP主：{meta['up']}　|　时长：{meta['duration'] // 60} 分 {meta['duration'] % 60} 秒"
        f"　|　BV号：{meta['bvid']}　|　字幕来源：{meta['subtitle_source']}"
    ).font.size = Pt(10)
    info.runs[0].font.color.rgb = GRAY
    set_cjk(info.runs[0].font)

    # ---------- 一图流（F14） ----------
    overview = next((f for f in c.get("figures", []) if f.get("role") == "overview"), None)
    other_figures = [f for f in c.get("figures", []) if f.get("role") != "overview"]
    if overview:
        doc.add_heading("一图流", level=1)
        picture(doc, video_dir, overview["file"], overview["caption"])

    # ---------- 视频总结版本 ----------
    doc.add_heading(c["summary"]["heading"], level=1)
    for f in other_figures:
        picture(doc, video_dir, f["file"], f["caption"])

    for i, ch in enumerate(c["summary"]["chapters"], 1):
        doc.add_heading(f"{i:02d} {ch['title']}（{ch['time_range']}）", level=2)
        doc.add_paragraph(ch["summary"])
        for b in ch["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
        for im in ch["images"]:
            picture(doc, video_dir, im["file"], f"{im['time']}　{im['caption']}")

    # ---------- 视频全量文字版 ----------
    doc.add_heading(c["full"]["heading"], level=1)
    note = doc.add_paragraph()
    nr = note.add_run(f"来源：{meta['subtitle_source']}（B站导出），已做断句与明显听写纠错；"
                      "时间码为段落起点，可跳回原视频定位。")
    nr.font.size = Pt(9)
    nr.font.color.rgb = GRAY
    set_cjk(nr.font)

    for b in c["full"]["blocks"]:
        p = doc.add_paragraph()
        tr = p.add_run(f"【{b['time']}】")
        tr.font.bold = True
        tr.font.color.rgb = RGBColor(0x4C, 0x7D, 0xD8)
        p.add_run("　" + b["text"])
        for im in b["images"]:
            picture(doc, video_dir, im["file"], f"{im['time']}　{im['caption']}")

    out = video_dir / f"{meta['num']}.《{meta['title']}》.docx"
    doc.save(out)
    print(f"已生成：{out}")


if __name__ == "__main__":
    main()
